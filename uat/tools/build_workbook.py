"""Builds the UAT workbook from the Playwright results.

    python uat/tools/build_workbook.py

Reads  uat/results/cycle-1/results.json   (first execution)
       uat/results/cycle-2/results.json   (re-test after fixes, optional)
       uat/defects.json                    (defect log)
       uat/tests/lib.mjs                   (test accounts)
Writes uat/NaijaPrimeSchool-UAT-Workbook.xlsx

Screenshots are embedded as thumbnails and hyperlinked to the full-size
images under uat/results/, so keep the workbook inside the uat folder.
"""
from __future__ import annotations

import io
import json
import re
import subprocess
from collections import OrderedDict
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from PIL import Image

UAT = Path(__file__).resolve().parents[1]
OUT = UAT / "NaijaPrimeSchool-UAT-Workbook.xlsx"

GREEN = "0B6B3A"
PASS_FILL = PatternFill("solid", fgColor="D8F0E0")
FAIL_FILL = PatternFill("solid", fgColor="F9D6D5")
HEAD_FILL = PatternFill("solid", fgColor=GREEN)
SUB_FILL = PatternFill("solid", fgColor="EAF4EE")
THIN = Side(style="thin", color="C9D3CC")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="top", wrap_text=True)
THUMB_W = 220


def load(path: Path):
    return json.loads(path.read_text(encoding="utf-8")) if path.exists() else None


def git(*args: str) -> str:
    try:
        return subprocess.check_output(["git", *args], cwd=UAT.parent, text=True).strip()
    except Exception:
        return ""


def header(ws, row: int, titles: list[str], widths: list[int] | None = None):
    for i, title in enumerate(titles, start=1):
        c = ws.cell(row=row, column=i, value=title)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
        c.alignment = CENTER
        c.border = BOX
    if widths:
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def status_fill(cell):
    v = str(cell.value or "")
    if v.startswith("Pass"):
        cell.fill = PASS_FILL
    elif v.startswith("Fail"):
        cell.fill = FAIL_FILL


def thumbnail(src: Path) -> io.BytesIO | None:
    if not src.exists():
        return None
    img = Image.open(src)
    # Keep the top of the page (header + main content), scaled to THUMB_W.
    ratio = THUMB_W / img.width
    crop_h = min(img.height, int(img.width * 0.68))
    img = img.crop((0, 0, img.width, crop_h)).resize((THUMB_W, int(crop_h * ratio)))
    buf = io.BytesIO()
    img.convert("RGB").save(buf, "JPEG", quality=70)
    buf.seek(0)
    return buf


def sprint_list(labels) -> str:
    """'Sprint 3', 'Sprints 3 & 9', 'Sprint 5b' ... -> 'Sprints 3, 5b, 9'."""
    labels = list(labels)
    ranges = [l for l in labels if "–" in l]
    if ranges:
        return ranges[0]
    nums = {n for l in labels for n in re.findall(r"\d+b?", l)}
    ordered = sorted(nums, key=lambda n: (int(n.rstrip("b")), n))
    return ("Sprint " if len(ordered) == 1 else "Sprints ") + ", ".join(ordered)


def accounts():
    text = (UAT / "tests" / "lib.mjs").read_text(encoding="utf-8")
    block = text[text.index("export const USERS"):text.index("};", text.index("export const USERS"))]
    rows = []
    for m in re.finditer(r"'?([\w.]+)'?: \{ password: '([^']+)', role: '([^']+)', name: '([^']+)' \}", block):
        rows.append(m.groups())
    return rows


def main():
    c1 = load(UAT / "results" / "cycle-1" / "results.json")
    c2 = load(UAT / "results" / "cycle-2" / "results.json")
    defects = load(UAT / "defects.json") or []
    if not c1:
        raise SystemExit("Run the UAT first: uat/results/cycle-1/results.json is missing.")

    r1 = OrderedDict((c["id"], c) for c in c1["cases"])
    r2 = {c["id"]: c for c in (c2 or {}).get("cases", [])}
    defect_by_test: dict[str, list[str]] = {}
    for d in defects:
        for t in d.get("tests", []):
            defect_by_test.setdefault(t, []).append(d["id"])

    def final(tid):
        return (r2.get(tid) or r1[tid])

    modules = OrderedDict()
    for c in r1.values():
        modules.setdefault(c["module"], []).append(c["id"])

    wb = Workbook()

    # ------------------------------------------------------------ summary sheet
    ws = wb.active
    ws.title = "Summary & Sign-off"
    ws.sheet_view.showGridLines = False
    ws["A1"] = "Naija Prime School — User Acceptance Test (UAT) Report"
    ws["A1"].font = Font(size=16, bold=True, color=GREEN)
    info = [
        ("System under test", "Naija Prime School Management Portal (Sprints 1–16)"),
        ("Build", f"{git('rev-parse', '--abbrev-ref', 'HEAD')} @ {git('rev-parse', '--short', 'HEAD')}"),
        ("Environment", f"{c1.get('baseUrl')} — ASP.NET Core (.NET 10), SQL Server 2022, {c1.get('browser')}"),
        ("Test data", "Demo school seeded by uat/seed (see the Test Data sheet)"),
        ("Execution", "Scripted UAT run in a real browser (Playwright), each case signed in as its role, with a screenshot per case"),
        ("Cycle 1 (first run)", c1.get("runAt", "")[:16].replace("T", " ") + " UTC"),
        ("Cycle 2 (re-test after fixes)", (c2 or {}).get("runAt", "not run")[:16].replace("T", " ") + (" UTC" if c2 else "")),
        ("Tester", "Claude (AI assistant) — automated execution; business review and sign-off below"),
    ]
    for i, (k, v) in enumerate(info, start=3):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True)
        ws.cell(row=i, column=2, value=v)

    row = len(info) + 5
    ws.cell(row=row, column=1, value="Results by module").font = Font(size=13, bold=True, color=GREEN)
    row += 1
    header(ws, row, ["Module", "Sprints", "Test cases", "Cycle 1 pass", "Cycle 1 fail", "Final pass", "Final fail", "Final pass rate"],
           [34, 22, 12, 13, 13, 12, 12, 15])
    totals = [0, 0, 0, 0, 0]
    for m, ids in modules.items():
        row += 1
        sprints = sprint_list(r1[i]["sprint"] for i in ids)
        p1 = sum(r1[i]["status"] == "Pass" for i in ids)
        pf = sum(final(i)["status"] == "Pass" for i in ids)
        vals = [m, sprints, len(ids), p1, len(ids) - p1, pf, len(ids) - pf, pf / len(ids)]
        for j, v in enumerate(vals, start=1):
            c = ws.cell(row=row, column=j, value=v)
            c.border = BOX
            c.alignment = WRAP if j <= 2 else CENTER
        ws.cell(row=row, column=8).number_format = "0%"
        for k, v in enumerate([len(ids), p1, len(ids) - p1, pf, len(ids) - pf]):
            totals[k] += v
    row += 1
    for j, v in enumerate(["Total", "", *totals, totals[3] / totals[0]], start=1):
        c = ws.cell(row=row, column=j, value=v)
        c.font = Font(bold=True)
        c.fill = SUB_FILL
        c.border = BOX
        c.alignment = CENTER if j > 2 else WRAP
    ws.cell(row=row, column=8).number_format = "0%"

    row += 2
    ws.cell(row=row, column=1, value="Defects").font = Font(size=13, bold=True, color=GREEN)
    row += 1
    header(ws, row, ["Severity", "Raised", "Fixed & verified", "Open"])
    for sev in ["Critical", "High", "Medium", "Low"]:
        row += 1
        ds = [d for d in defects if d["severity"] == sev]
        vals = [sev, len(ds), sum(d["status"] == "Fixed & verified" for d in ds), sum(d["status"] != "Fixed & verified" for d in ds)]
        for j, v in enumerate(vals, start=1):
            c = ws.cell(row=row, column=j, value=v)
            c.border = BOX
            c.alignment = CENTER

    row += 2
    ws.cell(row=row, column=1, value="Acceptance sign-off").font = Font(size=13, bold=True, color=GREEN)
    row += 1
    ws.cell(row=row, column=1, value="Each business owner reviews the module sheets (column \"Reviewer verdict\") and records a decision here.").alignment = WRAP
    row += 1
    header(ws, row, ["Role", "Name", "Modules reviewed", "Decision", "Conditions / comments", "Signature", "Date"])
    decision = DataValidation(type="list", formula1='"Accepted,Accepted with conditions,Rejected"', allow_blank=True)
    ws.add_data_validation(decision)
    for role, mods in [
        ("Project sponsor / Proprietor", "All"),
        ("Head teacher", "Academics, Students & parents, Attendance, Results, Announcements, Messaging, Audit log"),
        ("School bursar", "Fees & payments, Online payments"),
        ("Store keeper", "Store & inventory"),
        ("Class teacher representative", "Attendance, Results & report cards"),
        ("Parent representative", "Family portals, Announcements, Messaging, Online payments"),
        ("System administrator", "Sign-in & access control, User management, Audit log"),
    ]:
        row += 1
        for j, v in enumerate([role, "", mods, "", "", "", ""], start=1):
            c = ws.cell(row=row, column=j, value=v)
            c.border = BOX
            c.alignment = WRAP
        decision.add(ws.cell(row=row, column=4))
        ws.row_dimensions[row].height = 30
    for col, w in zip("ABCDEFG", [34, 22, 40, 22, 36, 18, 12]):
        ws.column_dimensions[col].width = max(ws.column_dimensions[col].width or 0, w)

    # -------------------------------------------------------------- test data
    td = wb.create_sheet("Test Data")
    td.sheet_view.showGridLines = False
    td["A1"] = "Test accounts (demo data only — reset with uat/reset-and-start.sh)"
    td["A1"].font = Font(size=13, bold=True, color=GREEN)
    header(td, 3, ["Username", "Password", "Role", "Who"], [22, 14, 20, 70])
    for i, (u, p, r, n) in enumerate(accounts(), start=4):
        for j, v in enumerate([u, p, r, n], start=1):
            c = td.cell(row=i, column=j, value=v)
            c.border = BOX
    base = len(accounts()) + 6
    facts = [
        "All other staff: password Uat@12345 (aisha.bello, emeka.nwosu, grace.udoh are teachers).",
        "All other parents and pupils: username firstname.lastname (e.g. zainab.bello), password Family@123.",
        "Current session and term: the one containing the run date; First Term fees for Primary 5 (₦257,000) and Primary 1 (₦222,000).",
        "30 pupils in 7 classes; Primary 5A (13 pupils, class teacher Mr. Chinedu Okeke) has the full timetable, registers, results and report cards.",
        "21 invoices: 12 paid, 6 half-paid, 3 unpaid (Kamsi Eze, Oluwaseun Akinola, Somto Obi).",
        "Store: 10 items; HB pencils and Whiteboard markers start below reorder level.",
        "Email and SMS use the Log providers (written to uat/results/app.log); online payments use the built-in Paystack simulator.",
    ]
    td.cell(row=base, column=1, value="Key facts").font = Font(bold=True)
    for i, f in enumerate(facts, start=base + 1):
        td.cell(row=i, column=1, value="• " + f)

    # ----------------------------------------------------------- module sheets
    verdict = DataValidation(type="list", formula1='"Accept,Reject,Query"', allow_blank=True)
    cols = ["Test ID", "Sprint", "Feature", "Role", "Signed in as", "Test case", "Preconditions", "Steps", "Expected result",
            "Cycle 1 actual result", "Cycle 1", "Re-test actual result", "Final status", "Defect(s)", "Tester", "Executed", "Evidence", "Reviewer verdict", "Reviewer comments"]
    widths = [10, 11, 16, 13, 16, 30, 24, 46, 42, 42, 9, 42, 10, 11, 12, 16, THUMB_W // 7 + 2, 12, 26]
    for m, ids in modules.items():
        sh = wb.create_sheet(re.sub(r"[\\/*?:\[\]]", "", m)[:31])
        sh.sheet_view.showGridLines = False
        dv = DataValidation(type="list", formula1='"Accept,Reject,Query"', allow_blank=True)
        sh.add_data_validation(dv)
        header(sh, 1, cols, widths)
        for i, tid in enumerate(ids, start=2):
            a, b = r1[tid], r2.get(tid)
            f = final(tid)
            steps = "\n".join(f"{n}. {s}" for n, s in enumerate(a["steps"], start=1))
            vals = [tid, a["sprint"], a["feature"], a["role"], a["user"], a["title"], a.get("preconditions") or "", steps, a["expected"],
                    a["actual"], a["status"], (b or {}).get("actual", "—" if not b else ""), f["status"],
                    ", ".join(defect_by_test.get(tid, [])), "Claude (automated)",
                    f["executedAt"][:16].replace("T", " "), "", "", ""]
            for j, v in enumerate(vals, start=1):
                c = sh.cell(row=i, column=j, value=v)
                c.alignment = CENTER if j in (1, 2, 11, 13, 14, 18) else WRAP
                c.border = BOX
            status_fill(sh.cell(row=i, column=11))
            status_fill(sh.cell(row=i, column=13))
            dv.add(sh.cell(row=i, column=18))

            shot_dir = "cycle-2" if b else "cycle-1"
            shot = UAT / "results" / shot_dir / f["screenshot"]
            thumb = thumbnail(shot)
            ev = sh.cell(row=i, column=17)
            if thumb:
                img = XLImage(thumb)
                sh.add_image(img, ev.coordinate)
                ev.hyperlink = f"results/{shot_dir}/{f['screenshot']}"
                ev.value = " "
                sh.row_dimensions[i].height = max(160, img.height * 0.75 + 6)
            else:
                sh.row_dimensions[i].height = 160
        sh.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(ids) + 1}"

    # ---------------------------------------------------------------- defects
    dl = wb.create_sheet("Defect Log")
    dl.sheet_view.showGridLines = False
    header(dl, 1, ["Defect ID", "Title", "Severity", "Module", "Sprint", "Found by test(s)", "Steps to reproduce", "Expected", "Actual (before fix)",
                   "Root cause", "Fix", "Status", "Re-test"], [10, 34, 10, 18, 10, 14, 40, 34, 34, 44, 44, 16, 30])
    for i, d in enumerate(defects, start=2):
        vals = [d["id"], d["title"], d["severity"], d["module"], d["sprint"], ", ".join(d["tests"]), d["steps"], d["expected"], d["actual"],
                d["cause"], d["fix"], d["status"], d.get("retest", "")]
        for j, v in enumerate(vals, start=1):
            c = dl.cell(row=i, column=j, value=v)
            c.alignment = WRAP
            c.border = BOX
        dl.cell(row=i, column=12).fill = PASS_FILL if d["status"] == "Fixed & verified" else FAIL_FILL
        dl.row_dimensions[i].height = 120
    dl.auto_filter.ref = f"A1:M{len(defects) + 1}"

    wb.save(OUT)
    print(f"Wrote {OUT.relative_to(UAT.parent)}: {len(r1)} test cases in {len(modules)} module sheets, {len(defects)} defects.")


if __name__ == "__main__":
    main()
