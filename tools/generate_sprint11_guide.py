"""Generates 'Sprint 11 - Implementation Guide.docx' covering email / SMS
notifications for unread announcements.

The prose lives in the markdown companion ('Sprint 11 - Implementation
Guide.md'); this script renders it with the same styling helpers as the
sprint 10 guide and appends full listings of the new source files, so the
two documents never drift apart.

Run from the repo root:  python tools/generate_sprint11_guide.py
"""

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from generate_sprint10_guide import (
    ROOT, add_bullets, add_code, add_file, add_heading, add_numbered,
    add_page_break, configure_document, set_cell_background, toc_page,
)

OUTPUT = "Sprint 11 - Implementation Guide.docx"
SOURCE_MD = "Sprint 11 - Implementation Guide.md"

LISTINGS = [
    "src/NaijaPrimeSchool.Domain/Communications/AnnouncementNotification.cs",
    "src/NaijaPrimeSchool.Domain/Communications/NotificationChannel.cs",
    "src/NaijaPrimeSchool.Domain/Communications/NotificationStatus.cs",
    "src/NaijaPrimeSchool.Application/Communications/INotificationService.cs",
    "src/NaijaPrimeSchool.Application/Communications/IMessageGateways.cs",
    "src/NaijaPrimeSchool.Application/Communications/Dtos/NotificationDtos.cs",
    "src/NaijaPrimeSchool.Infrastructure/Services/NotificationService.cs",
    "src/NaijaPrimeSchool.Infrastructure/Notifications/NotificationDispatcher.cs",
    "src/NaijaPrimeSchool.Infrastructure/Notifications/NotificationDispatchWorker.cs",
    "src/NaijaPrimeSchool.Infrastructure/Notifications/NotificationComposer.cs",
    "src/NaijaPrimeSchool.Infrastructure/Notifications/ContactNormalizer.cs",
    "src/NaijaPrimeSchool.Infrastructure/Notifications/NotificationOptions.cs",
    "src/NaijaPrimeSchool.Infrastructure/Notifications/LogGateways.cs",
    "src/NaijaPrimeSchool.Infrastructure/Notifications/SmtpEmailGateway.cs",
    "src/NaijaPrimeSchool.Infrastructure/Notifications/TermiiSmsGateway.cs",
    "src/NaijaPrimeSchool.Web/Components/Pages/Communications/NotificationLog.razor",
]

INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*|\[[^\]]+\]\([^)]+\))")


def add_rich(paragraph, text):
    """Renders **bold**, *italic*, `code` and [links](url) inline."""
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(0x05, 0x61, 0x3C)
        elif part.startswith("["):
            label, url = re.match(r"\[([^\]]+)\]\(([^)]+)\)", part).groups()
            paragraph.add_run(f"{label} ({url})")
        elif part.startswith("*"):
            paragraph.add_run(part[1:-1]).italic = True
        else:
            paragraph.add_run(part)


def add_table(doc, rows):
    header, body = rows[0], rows[1:]
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    for i, cell_text in enumerate(header):
        cell = table.rows[0].cells[i]
        set_cell_background(cell, "05613C")
        cell.paragraphs[0].text = ""
        run = cell.paragraphs[0].add_run(cell_text.strip("` "))
        run.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    for row in body:
        cells = table.add_row().cells
        for i, cell_text in enumerate(row):
            cells[i].paragraphs[0].text = ""
            add_rich(cells[i].paragraphs[0], cell_text)
    doc.add_paragraph()


def render_markdown(doc, text):
    lines = text.splitlines()
    i = 0
    paragraph = []

    def flush():
        if paragraph:
            add_rich(doc.add_paragraph(), " ".join(s.strip() for s in paragraph))
            paragraph.clear()

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            flush()
            block = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            add_code(doc, "\n".join(block))
        elif stripped.startswith("#"):
            flush()
            level = len(stripped) - len(stripped.lstrip("#"))
            if level > 1:  # the H1 is the title page
                add_heading(doc, stripped.lstrip("# ").strip().replace("`", ""), level - 1)
        elif stripped.startswith("|"):
            flush()
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(set(c) <= set("-: ") for c in cells):
                    rows.append(cells)
                i += 1
            add_table(doc, rows)
            continue
        elif re.match(r"^(- |\d+\. )", stripped):
            flush()
            numbered = stripped[0].isdigit()
            items = []
            while i < len(lines) and lines[i].strip():
                item = lines[i].strip()
                if re.match(r"^(- |\d+\. )", item):
                    items.append(re.sub(r"^(- |\d+\. )", "", item))
                else:
                    items[-1] += " " + item
                i += 1
            for item in items:
                add_rich(doc.add_paragraph(style="List Number" if numbered else "List Bullet"), item)
            continue
        elif not stripped:
            flush()
        else:
            paragraph.append(line)
        i += 1
    flush()


def title_page(doc):
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t.paragraph_format.space_before = Pt(120)
    r = t.add_run("Naija Prime School")
    r.font.size = Pt(32); r.font.bold = True
    r.font.color.rgb = RGBColor(0x05, 0x61, 0x3C)

    sub = doc.add_paragraph(); sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rs = sub.add_run("Sprint 11 — Email / SMS Notifications for Unread Announcements")
    rs.font.size = Pt(18); rs.font.color.rgb = RGBColor(0xB8, 0x86, 0x0B)

    sub2 = doc.add_paragraph(); sub2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rs2 = sub2.add_run("Queue on publish · Unread grace period · Background dispatcher · SMTP & Termii")
    rs2.font.size = Pt(14); rs2.italic = True

    doc.add_paragraph(); doc.add_paragraph()

    meta = doc.add_paragraph(); meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run("Author: Benjamin Fadina").bold = True
    meta.add_run("\nBuilt on: Sprints 1–10 (identity through the parent-deletion hotfix)")
    meta.add_run("\nStack: .NET 10, Blazor Web App (Auto), EF Core 10, SQL Server, Radzen Blazor")
    meta.add_run("\nRepository: https://github.com/benjaminsqlserver/NaijaPrimeSchool")
    meta.add_run("\nLicence: MIT — see LICENSE at the repo root")
    add_page_break(doc)


def listings(doc):
    add_page_break(doc)
    add_heading(doc, "Appendix — Full listings", 1)
    for rel in LISTINGS:
        add_file(doc, rel)


def build():
    doc = Document()
    configure_document(doc)
    title_page(doc)
    toc_page(doc)
    render_markdown(doc, (ROOT / SOURCE_MD).read_text(encoding="utf-8"))
    listings(doc)
    out_path = ROOT / OUTPUT
    doc.save(str(out_path))
    print(f"Generated: {out_path}")


if __name__ == "__main__":
    build()
