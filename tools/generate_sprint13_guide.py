"""Generates 'Sprint 13 - Implementation Guide.docx' covering two-way
messaging between families and the school office.

The prose lives in 'Sprint 13 - Implementation Guide.md'; this script
renders it with the sprint 11 markdown renderer and appends full listings of
the new source files.

Run from the repo root:  python tools/generate_sprint13_guide.py
"""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from generate_sprint10_guide import ROOT, add_file, add_heading, add_page_break, configure_document, toc_page
from generate_sprint11_guide import render_markdown

OUTPUT = "Sprint 13 - Implementation Guide.docx"
SOURCE_MD = "Sprint 13 - Implementation Guide.md"

LISTINGS = [
    "src/NaijaPrimeSchool.Domain/Messaging/MessageThread.cs",
    "src/NaijaPrimeSchool.Domain/Messaging/ThreadMessage.cs",
    "src/NaijaPrimeSchool.Domain/Messaging/MessageThreadStatus.cs",
    "src/NaijaPrimeSchool.Application/Messaging/IMessagingService.cs",
    "src/NaijaPrimeSchool.Application/Messaging/Dtos/MessagingDtos.cs",
    "src/NaijaPrimeSchool.Infrastructure/Services/MessagingService.cs",
    "src/NaijaPrimeSchool.Web/Components/Shared/MessageThreadView.razor",
    "src/NaijaPrimeSchool.Web/Components/Shared/NewMessageForm.razor",
    "src/NaijaPrimeSchool.Web/Components/Pages/Messaging/Inbox.razor",
    "src/NaijaPrimeSchool.Web/Components/Pages/Portals/PortalMessages.razor",
]


def title_page(doc):
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t.paragraph_format.space_before = Pt(120)
    r = t.add_run("Naija Prime School")
    r.font.size = Pt(32); r.font.bold = True
    r.font.color.rgb = RGBColor(0x05, 0x61, 0x3C)

    sub = doc.add_paragraph(); sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rs = sub.add_run("Sprint 13 — Two-way Messaging")
    rs.font.size = Pt(18); rs.font.color.rgb = RGBColor(0xB8, 0x86, 0x0B)

    sub2 = doc.add_paragraph(); sub2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rs2 = sub2.add_run("Family ↔ school office conversations · Shared inbox · Live refresh")
    rs2.font.size = Pt(14); rs2.italic = True

    doc.add_paragraph(); doc.add_paragraph()

    meta = doc.add_paragraph(); meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run("Author: Benjamin Fadina").bold = True
    meta.add_run("\nBuilt on: Sprints 1–12 (through notification preferences)")
    meta.add_run("\nStack: .NET 10, Blazor Web App (Auto), EF Core 10, SQL Server, Radzen Blazor")
    meta.add_run("\nRepository: https://github.com/benjaminsqlserver/NaijaPrimeSchool")
    meta.add_run("\nLicence: MIT — see LICENSE at the repo root")
    add_page_break(doc)


def build():
    doc = Document()
    configure_document(doc)
    title_page(doc)
    toc_page(doc)
    render_markdown(doc, (ROOT / SOURCE_MD).read_text(encoding="utf-8"))
    add_page_break(doc)
    add_heading(doc, "Appendix — Full listings", 1)
    for rel in LISTINGS:
        add_file(doc, rel)
    out_path = ROOT / OUTPUT
    doc.save(str(out_path))
    print(f"Generated: {out_path}")


if __name__ == "__main__":
    build()
